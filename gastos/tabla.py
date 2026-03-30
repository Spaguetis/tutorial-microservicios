import json
import sys
from datetime import date, datetime
from pathlib import Path
import tkinter as tk
from tkinter import messagebox, ttk

from docx import Document


def _base_dir() -> Path:
	"""Devuelve la carpeta del ejecutable o del script, segun corresponda."""
	if getattr(sys, "frozen", False):
		return Path(sys.executable).parent
	return Path(__file__).parent


DATA_FILE = _base_dir() / "gastos_data.json"
SYSTEM_MONTH_CHECK_MS = 60_000

# Emojis para categorías
CATEGORY_EMOJIS = {
	"Alimentacion": "🍔",
	"Transporte": "🚗",
	"Servicios": "💡",
	"Vivienda": "🏠",
	"Salud": "⚕️",
	"Entretenimiento": "🎬",
	"Ahorro": "💰",
	"Otros": "📦",
}

# Emojis para metas de ahorro
GOAL_EMOJIS = {
	"Emergencia": "🆘",
	"Viaje": "✈️",
	"Inversion": "📈",
	"Estudios": "🎓",
	"Hogar": "🏡",
	"Otro": "📌",
}

CATEGORIES = [
	"Alimentacion",
	"Transporte",
	"Servicios",
	"Vivienda",
	"Salud",
	"Entretenimiento",
	"Ahorro",
	"Otros",
]
SAVINGS_CURRENCIES = ["CLP", "USD"]
SAVINGS_GOALS = [
	"Emergencia",
	"Viaje",
	"Inversion",
	"Estudios",
	"Hogar",
	"Otro",
]

# Paleta de colores verde/azul profesional
COLORS = {
	"primary": "#1e7a4e",      # Verde oscuro profesional
	"primary_light": "#2d9d6a", # Verde más claro
	"secondary": "#0d7a7e",    # Azul-verde
	"accent": "#4cb85c",       # Verde vibrante
	"background": "#f0f8f5",   # Fondo muy claro
	"surface": "#ffffff",      # Blanco puro
	"text_primary": "#1a1a1a",  # Casi negro
	"text_secondary": "#666666", # Gris medio
	"success": "#28a745",      # Verde éxito
	"warning": "#ffc107",      # Amarillo
	"danger": "#dc3545",       # Rojo
	"info": "#17a2b8",         # Cian
}

CHART_COLORS = [
	"#1e7a4e",
	"#2d9d6a",
	"#0d7a7e",
	"#4cb85c",
	"#28a745",
	"#17a2b8",
	"#6f42c1",
	"#e83e8c",
]


def parse_amount(value: str) -> float:
	normalized = value.strip().replace(" ", "").replace(",", ".")
	amount = float(normalized)
	if amount < 0:
		raise ValueError("El monto no puede ser negativo")
	return amount


def validate_month(value: str) -> str:
	datetime.strptime(value, "%Y-%m")
	return value


def validate_date(value: str) -> str:
	datetime.strptime(value, "%Y-%m-%d")
	return value


def format_chilean_number(amount: float, decimals: int = 0) -> str:
	formatted = f"{amount:,.{decimals}f}"
	formatted = formatted.replace(",", "_").replace(".", ",").replace("_", ".")
	return formatted


def month_has_activity(data: dict, month: str) -> bool:
	if float(data.get("salaries", {}).get(month, 0.0)) > 0:
		return True
	if any(expense["date"].startswith(month) for expense in data.get("expenses", [])):
		return True
	return any(saving["date"].startswith(month) for saving in data.get("savings", []))


class ExpenseApp:
	def __init__(self, root: tk.Tk) -> None:
		self.root = root
		self.root.title("Control de gastos mensuales")
		self.root.geometry("1240x960")
		self.root.minsize(1120, 860)

		self.data = self.load_data()
		self.editing_expense_id = None
		self.editing_saving_id = None
		self.status_var = tk.StringVar(value="Listo")

		today = date.today()
		self.selected_month = tk.StringVar(value=today.strftime("%Y-%m"))
		self.salary_var = tk.StringVar()
		self.usd_rate_var = tk.StringVar()

		self.description_var = tk.StringVar()
		self.category_var = tk.StringVar(value=CATEGORIES[0])
		self.amount_var = tk.StringVar()
		self.expense_date_var = tk.StringVar(value=today.isoformat())

		self.saving_description_var = tk.StringVar()
		self.saving_goal_var = tk.StringVar(value=SAVINGS_GOALS[0])
		self.saving_currency_var = tk.StringVar(value=SAVINGS_CURRENCIES[0])
		self.saving_amount_var = tk.StringVar()
		self.saving_date_var = tk.StringVar(value=today.isoformat())

		self.current_salary_value = tk.StringVar()
		self.current_usd_rate_value = tk.StringVar()
		self.total_expenses_value = tk.StringVar()
		self.available_balance_value = tk.StringVar()
		self.expense_count_value = tk.StringVar()
		self.average_expense_value = tk.StringVar()
		self.savings_clp_value = tk.StringVar()
		self.savings_usd_value = tk.StringVar()
		self.savings_usd_clp_value = tk.StringVar()

		self.configure_style()
		self.build_ui()
		self.handle_system_month_change(show_message=False)
		self.refresh_view()
		self.schedule_system_month_check()

	def configure_style(self) -> None:
		style = ttk.Style()
		if "clam" in style.theme_names():
			style.theme_use("clam")

		# Colores para el tema
		bg_color = COLORS["background"]
		surface_color = COLORS["surface"]
		primary = COLORS["primary"]
		primary_light = COLORS["primary_light"]
		text_primary = COLORS["text_primary"]
		text_secondary = COLORS["text_secondary"]

		# Configurar colores del theme
		style.configure("TFrame", background=bg_color)
		style.configure("TLabelFrame", background=bg_color, foreground=text_primary)
		style.configure("TLabel", background=bg_color, foreground=text_primary)
		style.configure("TEntry", fieldbackground=surface_color, foreground=text_primary)
		style.configure("TCombobox", fieldbackground=surface_color, foreground=text_primary)
		style.configure("Treeview", background=surface_color, foreground=text_primary, fieldbackground=surface_color)
		style.configure("Treeview.Heading", background=primary, foreground=surface_color)
		style.map("Treeview.Heading", background=[("active", primary_light)])

		# Estilos personalizados
		style.configure("Title.TLabel", font=("Segoe UI", 24, "bold"), foreground=primary)
		style.configure("Heading.TLabel", font=("Segoe UI", 14, "bold"), foreground=primary)
		style.configure("CardValue.TLabel", font=("Segoe UI", 13, "bold"), foreground=primary)
		style.configure("Muted.TLabel", foreground=text_secondary, font=("Segoe UI", 10))
		style.configure("MonoValue.TLabel", font=("Courier New", 14, "bold"), foreground=primary)

		# Estilos para botones
		style.configure("TButton", font=("Segoe UI", 10), padding=6)
		style.map("TButton",
			background=[("active", primary_light), ("pressed", primary)],
			foreground=[("active", surface_color)]
		)

		# Entrada de texto mejorada
		style.configure("TLabelFrame.Label", font=("Segoe UI", 11, "bold"), foreground=primary)

	def default_data(self) -> dict:
		return {
			"salaries": {},
			"usd_rates": {},
			"expenses": [],
			"savings": [],
			"metadata": {"last_seen_system_month": date.today().strftime("%Y-%m")},
		}

	def load_data(self) -> dict:
		if not DATA_FILE.exists():
			return self.default_data()

		try:
			with DATA_FILE.open("r", encoding="utf-8") as file:
				data = json.load(file)
		except (json.JSONDecodeError, OSError):
			messagebox.showwarning(
				"Datos invalidos",
				"No se pudo leer el archivo de datos. Se iniciara una base vacia.",
			)
			return self.default_data()

		data.setdefault("salaries", {})
		data.setdefault("usd_rates", {})
		data.setdefault("expenses", [])
		data.setdefault("savings", [])
		data.setdefault("metadata", {})
		data["metadata"].setdefault("last_seen_system_month", date.today().strftime("%Y-%m"))
		return data

	def save_data(self) -> None:
		self.data.setdefault("metadata", {})
		with DATA_FILE.open("w", encoding="utf-8") as file:
			json.dump(self.data, file, indent=2, ensure_ascii=False)

	def build_ui(self) -> None:
		container = ttk.Frame(self.root, padding=16)
		container.pack(fill="both", expand=True)

		# Header
		header_frame = ttk.Frame(container)
		header_frame.pack(fill="x", pady=(0, 20))

		title = ttk.Label(container, text="💰 Gestor de Gastos y Ahorros", style="Title.TLabel")
		title.pack(anchor="w")

		subtitle = ttk.Label(
			container,
			text="Controla tus gastos en CLP, ahorros en CLP/USD y metas financieras de forma profesional.",
			style="Muted.TLabel",
		)
		subtitle.pack(anchor="w", pady=(4, 16))

		top_frame = ttk.Frame(container)
		top_frame.pack(fill="x")

		self.build_month_section(top_frame)
		self.build_salary_section(top_frame)
		self.build_exchange_rate_section(top_frame)

		middle_frame = ttk.Frame(container)
		middle_frame.pack(fill="x", pady=(16, 12))

		self.build_expense_form(middle_frame)
		self.build_savings_form(middle_frame)
		self.build_summary_section(middle_frame)

		tables_frame = ttk.Frame(container)
		tables_frame.pack(fill="both", expand=True)
		tables_frame.columnconfigure(0, weight=3)
		tables_frame.columnconfigure(1, weight=2)
		tables_frame.rowconfigure(0, weight=3)
		tables_frame.rowconfigure(1, weight=2)
		tables_frame.rowconfigure(2, weight=2)

		self.build_expense_table(tables_frame)
		self.build_category_table(tables_frame)
		self.build_savings_table(tables_frame)
		self.build_chart_section(tables_frame)

		status_bar = ttk.Label(container, textvariable=self.status_var, style="Muted.TLabel")
		status_bar.pack(fill="x", pady=(10, 0))
		self.root.configure(bg=COLORS["background"])

	def build_month_section(self, parent: ttk.Frame) -> None:
		frame = ttk.LabelFrame(parent, text="📅 Mes de trabajo", padding=12)
		frame.pack(side="left", fill="x", expand=True, padx=(0, 8))

		ttk.Label(frame, text="Mes (YYYY-MM)", style="Muted.TLabel").grid(row=0, column=0, sticky="w")
		month_entry = ttk.Entry(frame, textvariable=self.selected_month, width=14)
		month_entry.grid(row=1, column=0, padx=(0, 8), pady=(4, 0), sticky="w")

		ttk.Button(frame, text="📥 Cargar", command=self.change_month).grid(row=1, column=1, sticky="w", pady=(4, 0))
		ttk.Button(frame, text="📍 Hoy", command=self.go_to_current_month).grid(row=1, column=2, padx=(8, 0), sticky="w", pady=(4, 0))
		ttk.Button(frame, text="📄 Exportar", command=self.export_selected_month_report).grid(row=1, column=3, padx=(8, 0), sticky="w", pady=(4, 0))

	def build_salary_section(self, parent: ttk.Frame) -> None:
		frame = ttk.LabelFrame(parent, text="💼 Salario mensual en CLP", padding=12)
		frame.pack(side="left", fill="x", expand=True, padx=(0, 8))

		ttk.Label(frame, text="Monto del salario", style="Muted.TLabel").grid(row=0, column=0, sticky="w")
		ttk.Entry(frame, textvariable=self.salary_var, width=16).grid(row=1, column=0, padx=(0, 8), pady=(4, 0), sticky="w")
		ttk.Button(frame, text="💾 Guardar", command=self.save_salary).grid(row=1, column=1, sticky="w", pady=(4, 0))

		ttk.Label(frame, text="Cargado:", style="Muted.TLabel").grid(row=2, column=0, columnspan=2, sticky="w", pady=(10, 0))
		ttk.Label(frame, textvariable=self.current_salary_value, style="CardValue.TLabel").grid(row=3, column=0, columnspan=2, sticky="w")

	def build_exchange_rate_section(self, parent: ttk.Frame) -> None:
		frame = ttk.LabelFrame(parent, text="💱 Tipo de cambio USD", padding=12)
		frame.pack(side="left", fill="x", expand=True)

		ttk.Label(frame, text="Valor 1 USD en CLP", style="Muted.TLabel").grid(row=0, column=0, sticky="w")
		ttk.Entry(frame, textvariable=self.usd_rate_var, width=16).grid(row=1, column=0, padx=(0, 8), pady=(4, 0), sticky="w")
		ttk.Button(frame, text="💾 Guardar", command=self.save_usd_rate).grid(row=1, column=1, sticky="w", pady=(4, 0))

		ttk.Label(frame, text="Cargado:", style="Muted.TLabel").grid(row=2, column=0, columnspan=2, sticky="w", pady=(10, 0))
		ttk.Label(frame, textvariable=self.current_usd_rate_value, style="CardValue.TLabel").grid(row=3, column=0, columnspan=2, sticky="w")

	def build_expense_form(self, parent: ttk.Frame) -> None:
		frame = ttk.LabelFrame(parent, text="📝 Registrar gasto en CLP", padding=12)
		frame.pack(side="left", fill="both", expand=True, padx=(0, 8))

		ttk.Label(frame, text="Descripción", style="Muted.TLabel").grid(row=0, column=0, sticky="w")
		ttk.Entry(frame, textvariable=self.description_var, width=26).grid(row=1, column=0, padx=(0, 8), pady=(4, 10), sticky="ew")

		ttk.Label(frame, text="Categoría", style="Muted.TLabel").grid(row=0, column=1, sticky="w")
		category_options = [f"{CATEGORY_EMOJIS.get(cat, '•')} {cat}" for cat in CATEGORIES]
		self.category_combo_display = ttk.Combobox(frame, values=category_options, width=16, state="readonly")
		self.category_combo_display.grid(row=1, column=1, padx=(0, 8), pady=(4, 10), sticky="ew")
		self.category_combo_display.bind("<<ComboboxSelected>>", self._on_category_selected)
		if category_options:
			self.category_combo_display.set(category_options[0])

		ttk.Label(frame, text="Monto CLP", style="Muted.TLabel").grid(row=2, column=0, sticky="w")
		ttk.Entry(frame, textvariable=self.amount_var, width=18).grid(row=3, column=0, padx=(0, 8), pady=(4, 10), sticky="ew")

		ttk.Label(frame, text="Fecha (YYYY-MM-DD)", style="Muted.TLabel").grid(row=2, column=1, sticky="w")
		ttk.Entry(frame, textvariable=self.expense_date_var, width=18).grid(row=3, column=1, padx=(0, 8), pady=(4, 10), sticky="ew")

		buttons_frame = ttk.Frame(frame)
		buttons_frame.grid(row=4, column=0, columnspan=2, sticky="w")
		ttk.Button(buttons_frame, text="💾 Guardar gasto", command=self.submit_expense).pack(side="left")
		ttk.Button(buttons_frame, text="🗑️  Limpiar", command=self.reset_expense_form).pack(side="left", padx=(8, 0))

		frame.columnconfigure(0, weight=1)
		frame.columnconfigure(1, weight=1)

	def _on_category_selected(self, _event: object = None) -> None:
		selected_text = self.category_combo_display.get()
		if selected_text:
			category = selected_text.split(" ", 1)[1] if " " in selected_text else selected_text
			self.category_var.set(category)

	def build_savings_form(self, parent: ttk.Frame) -> None:
		frame = ttk.LabelFrame(parent, text="🎯 Registrar ahorro", padding=12)
		frame.pack(side="left", fill="both", expand=True, padx=(0, 8))

		ttk.Label(frame, text="Descripción", style="Muted.TLabel").grid(row=0, column=0, sticky="w")
		ttk.Entry(frame, textvariable=self.saving_description_var, width=26).grid(row=1, column=0, padx=(0, 8), pady=(4, 10), sticky="ew")

		ttk.Label(frame, text="Meta", style="Muted.TLabel").grid(row=0, column=1, sticky="w")
		goal_options = [f"{GOAL_EMOJIS.get(goal, '•')} {goal}" for goal in SAVINGS_GOALS]
		self.goal_combo_display = ttk.Combobox(frame, values=goal_options, width=14, state="readonly")
		self.goal_combo_display.grid(row=1, column=1, padx=(0, 8), pady=(4, 10), sticky="ew")
		self.goal_combo_display.bind("<<ComboboxSelected>>", self._on_goal_selected)
		if goal_options:
			self.goal_combo_display.set(goal_options[0])

		ttk.Label(frame, text="Moneda", style="Muted.TLabel").grid(row=2, column=0, sticky="w")
		currency_combo = ttk.Combobox(frame, textvariable=self.saving_currency_var, values=["🇨🇱 CLP", "🇺🇸 USD"], width=10, state="readonly")
		currency_combo.grid(row=3, column=0, padx=(0, 8), pady=(4, 10), sticky="ew")
		currency_combo.set("🇨🇱 CLP")

		ttk.Label(frame, text="Monto", style="Muted.TLabel").grid(row=2, column=1, sticky="w")
		ttk.Entry(frame, textvariable=self.saving_amount_var, width=18).grid(row=3, column=1, padx=(0, 8), pady=(4, 10), sticky="ew")

		ttk.Label(frame, text="Fecha (YYYY-MM-DD)", style="Muted.TLabel").grid(row=4, column=0, sticky="w")
		ttk.Entry(frame, textvariable=self.saving_date_var, width=18).grid(row=5, column=0, padx=(0, 8), pady=(4, 10), sticky="ew")

		buttons_frame = ttk.Frame(frame)
		buttons_frame.grid(row=5, column=1, sticky="w")
		ttk.Button(buttons_frame, text="💾 Guardar", command=self.submit_saving).pack(side="left")
		ttk.Button(buttons_frame, text="🗑️  Limpiar", command=self.reset_saving_form).pack(side="left", padx=(8, 0))

		frame.columnconfigure(0, weight=1)
		frame.columnconfigure(1, weight=1)

	def _on_goal_selected(self, _event: object = None) -> None:
		selected_text = self.goal_combo_display.get()
		if selected_text:
			goal = selected_text.split(" ", 1)[1] if " " in selected_text else selected_text
			self.saving_goal_var.set(goal)

	def build_summary_section(self, parent: ttk.Frame) -> None:
		frame = ttk.LabelFrame(parent, text="📊 Resumen del mes", padding=12)
		frame.pack(side="left", fill="both", expand=True)

		items = [
			("💼 Salario cargado", self.current_salary_value),
			("💱 Tipo de cambio USD", self.current_usd_rate_value),
			("💸 Total gastado", self.total_expenses_value),
			("💚 Ahorro en CLP", self.savings_clp_value),
			("🇺🇸 Ahorro en USD", self.savings_usd_value),
			("🔄 USD equivalente CLP", self.savings_usd_clp_value),
			("✅ Disponible final", self.available_balance_value),
			("📈 Cantidad de gastos", self.expense_count_value),
			("📊 Gasto promedio", self.average_expense_value),
		]

		for index, (label, variable) in enumerate(items):
			base_row = index * 2
			ttk.Label(frame, text=label, style="Muted.TLabel").grid(row=base_row, column=0, sticky="w")
			ttk.Label(frame, textvariable=variable, style="CardValue.TLabel").grid(row=base_row + 1, column=0, sticky="w", pady=(2, 6))

	def build_expense_table(self, parent: ttk.Frame) -> None:
		frame = ttk.LabelFrame(parent, text="📋 Gastos del mes en CLP", padding=10)
		frame.grid(row=0, column=0, sticky="nsew", padx=(0, 8), pady=(0, 8))
		frame.rowconfigure(0, weight=1)
		frame.columnconfigure(0, weight=1)

		columns = ("date", "description", "category", "amount")
		self.expense_tree = ttk.Treeview(frame, columns=columns, show="headings", height=12)
		self.expense_tree.heading("date", text="📅 Fecha")
		self.expense_tree.heading("description", text="📝 Descripción")
		self.expense_tree.heading("category", text="🏷️  Categoría")
		self.expense_tree.heading("amount", text="💰 Monto CLP")

		self.expense_tree.column("date", width=110, anchor="center")
		self.expense_tree.column("description", width=260)
		self.expense_tree.column("category", width=130)
		self.expense_tree.column("amount", width=130, anchor="e")

		self.expense_tree.grid(row=0, column=0, sticky="nsew")
		self.expense_tree.bind("<<TreeviewSelect>>", self.load_selected_expense_into_form)
		scrollbar = ttk.Scrollbar(frame, orient="vertical", command=self.expense_tree.yview)
		scrollbar.grid(row=0, column=1, sticky="ns")
		self.expense_tree.configure(yscrollcommand=scrollbar.set)

		buttons_frame = ttk.Frame(frame)
		buttons_frame.grid(row=1, column=0, sticky="w", pady=(10, 0))
		ttk.Button(buttons_frame, text="✏️  Editar", command=self.edit_selected_expense).pack(side="left")
		ttk.Button(buttons_frame, text="🗑️  Eliminar", command=self.delete_selected_expense).pack(side="left", padx=(8, 0))

	def build_category_table(self, parent: ttk.Frame) -> None:
		frame = ttk.LabelFrame(parent, text="📊 Totales por categoría", padding=10)
		frame.grid(row=0, column=1, sticky="nsew", pady=(0, 8))
		frame.rowconfigure(0, weight=1)
		frame.columnconfigure(0, weight=1)

		columns = ("category", "total")
		self.category_tree = ttk.Treeview(frame, columns=columns, show="headings", height=12)
		self.category_tree.heading("category", text="🏷️  Categoría")
		self.category_tree.heading("total", text="💰 Total CLP")
		self.category_tree.column("category", width=180)
		self.category_tree.column("total", width=130, anchor="e")
		self.category_tree.grid(row=0, column=0, sticky="nsew")

	def build_savings_table(self, parent: ttk.Frame) -> None:
		frame = ttk.LabelFrame(parent, text="💚 Ahorros del mes", padding=10)
		frame.grid(row=1, column=0, columnspan=2, sticky="nsew", pady=(0, 8))
		frame.rowconfigure(0, weight=1)
		frame.columnconfigure(0, weight=1)

		columns = ("date", "description", "goal", "currency", "amount", "clp_equivalent")
		self.savings_tree = ttk.Treeview(frame, columns=columns, show="headings", height=8)
		self.savings_tree.heading("date", text="📅 Fecha")
		self.savings_tree.heading("description", text="📝 Descripción")
		self.savings_tree.heading("goal", text="🎯 Meta")
		self.savings_tree.heading("currency", text="💱 Moneda")
		self.savings_tree.heading("amount", text="💰 Monto")
		self.savings_tree.heading("clp_equivalent", text="🔄 Equiv. CLP")

		self.savings_tree.column("date", width=110, anchor="center")
		self.savings_tree.column("description", width=220)
		self.savings_tree.column("goal", width=130)
		self.savings_tree.column("currency", width=90, anchor="center")
		self.savings_tree.column("amount", width=130, anchor="e")
		self.savings_tree.column("clp_equivalent", width=140, anchor="e")

		self.savings_tree.grid(row=0, column=0, sticky="nsew")
		self.savings_tree.bind("<<TreeviewSelect>>", self.load_selected_saving_into_form)
		scrollbar = ttk.Scrollbar(frame, orient="vertical", command=self.savings_tree.yview)
		scrollbar.grid(row=0, column=1, sticky="ns")
		self.savings_tree.configure(yscrollcommand=scrollbar.set)

		buttons_frame = ttk.Frame(frame)
		buttons_frame.grid(row=1, column=0, sticky="w", pady=(10, 0))
		ttk.Button(buttons_frame, text="✏️  Editar", command=self.edit_selected_saving).pack(side="left")
		ttk.Button(buttons_frame, text="🗑️  Eliminar", command=self.delete_selected_saving).pack(side="left", padx=(8, 0))

	def build_chart_section(self, parent: ttk.Frame) -> None:
		frame = ttk.Frame(parent)
		frame.grid(row=2, column=0, columnspan=2, sticky="nsew")
		frame.columnconfigure(0, weight=1)
		frame.columnconfigure(1, weight=1)
		frame.rowconfigure(0, weight=1)

		chart_frame = ttk.LabelFrame(frame, text="📉 Gastos por categoría", padding=10)
		chart_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 4))
		chart_frame.columnconfigure(0, weight=1)
		chart_frame.rowconfigure(0, weight=1)

		self.chart_canvas = tk.Canvas(chart_frame, background=COLORS["background"], highlightthickness=0, height=220)
		self.chart_canvas.grid(row=0, column=0, sticky="nsew")

		savings_frame = ttk.LabelFrame(frame, text="📈 Ahorros por meta", padding=10)
		savings_frame.grid(row=0, column=1, sticky="nsew", padx=(4, 0))
		savings_frame.columnconfigure(0, weight=1)
		savings_frame.rowconfigure(0, weight=1)

		self.savings_chart_canvas = tk.Canvas(savings_frame, background=COLORS["background"], highlightthickness=0, height=220)
		self.savings_chart_canvas.grid(row=0, column=0, sticky="nsew")

	def format_currency(self, amount: float, currency: str = "CLP") -> str:
		if currency == "USD":
			return f"US$ {format_chilean_number(amount, 2)}"
		return f"CLP$ {format_chilean_number(round(amount), 0)}"

	def set_status(self, message: str) -> None:
		self.status_var.set(message)

	def get_current_salary(self, month: str | None = None) -> float:
		month = month or self.selected_month.get()
		return float(self.data["salaries"].get(month, 0.0))

	def get_current_usd_rate(self, month: str | None = None) -> float:
		month = month or self.selected_month.get()
		return float(self.data["usd_rates"].get(month, 0.0))

	def get_month_expenses(self, month: str | None = None) -> list[dict]:
		month = month or self.selected_month.get()
		return [expense for expense in self.data["expenses"] if expense["date"].startswith(month)]

	def get_month_savings(self, month: str | None = None) -> list[dict]:
		month = month or self.selected_month.get()
		return [saving for saving in self.data["savings"] if saving["date"].startswith(month)]

	def get_savings_totals(self, month: str | None = None) -> dict[str, float]:
		usd_rate = self.get_current_usd_rate(month)
		totals = {
			"CLP": 0.0,
			"USD": 0.0,
			"USD_TO_CLP": 0.0,
			"TOTAL_CLP_EQUIVALENT": 0.0,
		}
		for saving in self.get_month_savings(month):
			currency = saving["currency"]
			amount = float(saving["amount"])
			if currency == "USD":
				totals["USD"] += amount
			else:
				totals["CLP"] += amount

		totals["USD_TO_CLP"] = totals["USD"] * usd_rate
		totals["TOTAL_CLP_EQUIVALENT"] = totals["CLP"] + totals["USD_TO_CLP"]
		return totals

	def get_savings_totals_by_goal(self, month: str | None = None) -> dict[str, float]:
		usd_rate = self.get_current_usd_rate(month)
		totals_by_goal: dict[str, float] = {}
		for saving in self.get_month_savings(month):
			goal = saving.get("goal", "Otro")
			amount = float(saving["amount"])
			if saving["currency"] == "USD":
				amount *= usd_rate
			totals_by_goal[goal] = totals_by_goal.get(goal, 0.0) + amount
		return totals_by_goal

	def go_to_current_month(self) -> None:
		self.selected_month.set(date.today().strftime("%Y-%m"))
		self.change_month()

	def change_month(self) -> None:
		try:
			month = validate_month(self.selected_month.get().strip())
		except ValueError:
			messagebox.showerror("Mes invalido", "Usa el formato YYYY-MM, por ejemplo 2026-03.")
			return

		self.selected_month.set(month)
		default_date = f"{month}-01"
		if not self.expense_date_var.get().startswith(month):
			self.expense_date_var.set(default_date)
		if not self.saving_date_var.get().startswith(month):
			self.saving_date_var.set(default_date)
		self.reset_expense_form(clear_date=False)
		self.reset_saving_form(clear_date=False)
		self.refresh_view()
		self.set_status(f"Mes cargado: {month}")

	def save_salary(self) -> None:
		try:
			month = validate_month(self.selected_month.get().strip())
			salary = parse_amount(self.salary_var.get())
		except ValueError:
			messagebox.showerror("Salario invalido", "Ingresa un mes valido y un salario numerico mayor o igual a 0.")
			return

		self.data["salaries"][month] = salary
		self.save_data()
		self.refresh_view()
		self.set_status(f"Salario actualizado para {month}")

	def save_usd_rate(self) -> None:
		try:
			month = validate_month(self.selected_month.get().strip())
			rate = parse_amount(self.usd_rate_var.get())
		except ValueError:
			messagebox.showerror("Tipo de cambio invalido", "Ingresa un mes valido y un tipo de cambio numerico mayor a 0.")
			return

		if rate == 0:
			messagebox.showerror("Tipo de cambio invalido", "El tipo de cambio debe ser mayor que 0.")
			return

		self.data["usd_rates"][month] = rate
		self.save_data()
		self.refresh_view()
		self.set_status(f"Tipo de cambio actualizado para {month}")

	def submit_expense(self) -> None:
		if self.editing_expense_id:
			self.update_expense()
			return
		self.add_expense()

	def submit_saving(self) -> None:
		if self.editing_saving_id:
			self.update_saving()
			return
		self.add_saving()

	def add_expense(self) -> None:
		expense = self.build_expense_from_form(str(int(datetime.now().timestamp() * 1000000)))
		if expense is None:
			return

		self.data["expenses"].append(expense)
		self.save_data()
		self.reset_expense_form()
		self.refresh_view()
		self.set_status("Gasto agregado")

	def update_expense(self) -> None:
		if not self.editing_expense_id:
			return

		updated_expense = self.build_expense_from_form(self.editing_expense_id)
		if updated_expense is None:
			return

		for index, expense in enumerate(self.data["expenses"]):
			if expense["id"] == self.editing_expense_id:
				self.data["expenses"][index] = updated_expense
				break

		self.save_data()
		self.reset_expense_form()
		self.refresh_view()
		self.set_status("Gasto actualizado")

	def add_saving(self) -> None:
		saving = self.build_saving_from_form(str(int(datetime.now().timestamp() * 1000000)))
		if saving is None:
			return

		self.data["savings"].append(saving)
		self.save_data()
		self.reset_saving_form()
		self.refresh_view()
		self.set_status("Ahorro agregado")

	def update_saving(self) -> None:
		if not self.editing_saving_id:
			return

		updated_saving = self.build_saving_from_form(self.editing_saving_id)
		if updated_saving is None:
			return

		for index, saving in enumerate(self.data["savings"]):
			if saving["id"] == self.editing_saving_id:
				self.data["savings"][index] = updated_saving
				break

		self.save_data()
		self.reset_saving_form()
		self.refresh_view()
		self.set_status("Ahorro actualizado")

	def build_expense_from_form(self, expense_id: str) -> dict | None:
		description = self.description_var.get().strip()
		if not description:
			messagebox.showerror("Descripcion requerida", "Escribe una descripcion para el gasto.")
			return None

		try:
			month = validate_month(self.selected_month.get().strip())
			expense_date = validate_date(self.expense_date_var.get().strip())
			amount = parse_amount(self.amount_var.get())
		except ValueError:
			messagebox.showerror("Datos invalidos", "Verifica el mes, la fecha y el monto. Usa YYYY-MM-DD para la fecha.")
			return None

		if not expense_date.startswith(month):
			messagebox.showerror("Fecha fuera del mes", "La fecha del gasto debe pertenecer al mes que estas gestionando.")
			return None

		return {
			"id": expense_id,
			"date": expense_date,
			"description": description,
			"category": self.category_var.get(),
			"amount": amount,
		}

	def build_saving_from_form(self, saving_id: str) -> dict | None:
		description = self.saving_description_var.get().strip()
		if not description:
			messagebox.showerror("Descripcion requerida", "Escribe una descripcion para el ahorro.")
			return None

		try:
			month = validate_month(self.selected_month.get().strip())
			saving_date = validate_date(self.saving_date_var.get().strip())
			amount = parse_amount(self.saving_amount_var.get())
		except ValueError:
			messagebox.showerror("Datos invalidos", "Verifica el mes, la fecha y el monto del ahorro.")
			return None

		if not saving_date.startswith(month):
			messagebox.showerror("Fecha fuera del mes", "La fecha del ahorro debe pertenecer al mes que estas gestionando.")
			return None

		goal = self.saving_goal_var.get()
		if goal not in SAVINGS_GOALS:
			messagebox.showerror("Meta invalida", "Selecciona una meta valida para el ahorro.")
			return None

		currency = self.saving_currency_var.get()
		if currency not in SAVINGS_CURRENCIES:
			messagebox.showerror("Moneda invalida", "Selecciona CLP o USD para el ahorro.")
			return None

		if currency == "USD" and self.get_current_usd_rate() == 0:
			messagebox.showerror("Falta tipo de cambio", "Antes de guardar ahorro en USD debes cargar el tipo de cambio del mes.")
			return None

		return {
			"id": saving_id,
			"date": saving_date,
			"description": description,
			"goal": goal,
			"currency": currency,
			"amount": amount,
		}

	def reset_expense_form(self, clear_date: bool = True) -> None:
		self.editing_expense_id = None
		self.description_var.set("")
		self.category_var.set(CATEGORIES[0])
		self.amount_var.set("")
		if clear_date:
			month = self.selected_month.get().strip()
			default_date = date.today().isoformat()
			if not default_date.startswith(month):
				default_date = f"{month}-01"
			self.expense_date_var.set(default_date)
		for item in self.expense_tree.selection():
			self.expense_tree.selection_remove(item)

	def reset_saving_form(self, clear_date: bool = True) -> None:
		self.editing_saving_id = None
		self.saving_description_var.set("")
		self.saving_goal_var.set(SAVINGS_GOALS[0])
		self.saving_currency_var.set(SAVINGS_CURRENCIES[0])
		self.saving_amount_var.set("")
		if clear_date:
			month = self.selected_month.get().strip()
			default_date = date.today().isoformat()
			if not default_date.startswith(month):
				default_date = f"{month}-01"
			self.saving_date_var.set(default_date)
		for item in self.savings_tree.selection():
			self.savings_tree.selection_remove(item)

	def load_selected_expense_into_form(self, _event: object = None) -> None:
		selected = self.expense_tree.selection()
		if not selected:
			return

		expense_id = selected[0]
		expense = next((item for item in self.data["expenses"] if item["id"] == expense_id), None)
		if expense is None:
			return

		self.editing_expense_id = expense_id
		self.description_var.set(expense["description"])
		self.category_var.set(expense["category"])
		self.amount_var.set(f"{expense['amount']:.2f}")
		self.expense_date_var.set(expense["date"])
		self.set_status("Edicion de gasto preparada")

	def load_selected_saving_into_form(self, _event: object = None) -> None:
		selected = self.savings_tree.selection()
		if not selected:
			return

		saving_id = selected[0]
		saving = next((item for item in self.data["savings"] if item["id"] == saving_id), None)
		if saving is None:
			return

		self.editing_saving_id = saving_id
		self.saving_description_var.set(saving["description"])
		self.saving_goal_var.set(saving.get("goal", SAVINGS_GOALS[0]))
		self.saving_currency_var.set(saving["currency"])
		self.saving_amount_var.set(f"{saving['amount']:.2f}")
		self.saving_date_var.set(saving["date"])
		self.set_status("Edicion de ahorro preparada")

	def edit_selected_expense(self) -> None:
		if not self.expense_tree.selection():
			messagebox.showinfo("Seleccion requerida", "Elige un gasto para editar.")
			return
		self.load_selected_expense_into_form()

	def edit_selected_saving(self) -> None:
		if not self.savings_tree.selection():
			messagebox.showinfo("Seleccion requerida", "Elige un ahorro para editar.")
			return
		self.load_selected_saving_into_form()

	def delete_selected_expense(self) -> None:
		selected = self.expense_tree.selection()
		if not selected:
			messagebox.showinfo("Seleccion requerida", "Elige un gasto para eliminar.")
			return

		expense_id = selected[0]
		self.data["expenses"] = [expense for expense in self.data["expenses"] if expense["id"] != expense_id]
		self.save_data()
		self.reset_expense_form()
		self.refresh_view()
		self.set_status("Gasto eliminado")

	def delete_selected_saving(self) -> None:
		selected = self.savings_tree.selection()
		if not selected:
			messagebox.showinfo("Seleccion requerida", "Elige un ahorro para eliminar.")
			return

		saving_id = selected[0]
		self.data["savings"] = [saving for saving in self.data["savings"] if saving["id"] != saving_id]
		self.save_data()
		self.reset_saving_form()
		self.refresh_view()
		self.set_status("Ahorro eliminado")

	def get_month_totals_by_category(self, month: str | None = None) -> dict[str, float]:
		totals_by_category: dict[str, float] = {}
		for expense in self.get_month_expenses(month):
			category = expense["category"]
			totals_by_category[category] = totals_by_category.get(category, 0.0) + float(expense["amount"])
		return totals_by_category

	def refresh_view(self) -> None:
		current_salary = self.get_current_salary()
		usd_rate = self.get_current_usd_rate()
		month_expenses = self.get_month_expenses()
		total_expenses = sum(float(expense["amount"]) for expense in month_expenses)
		expense_count = len(month_expenses)
		average_expense = total_expenses / expense_count if expense_count else 0.0
		savings_totals = self.get_savings_totals()
		available_balance = current_salary - total_expenses - savings_totals["TOTAL_CLP_EQUIVALENT"]

		self.salary_var.set(f"{current_salary:.2f}" if current_salary else "")
		self.usd_rate_var.set(f"{usd_rate:.2f}" if usd_rate else "")
		self.current_salary_value.set(self.format_currency(current_salary, "CLP"))
		self.current_usd_rate_value.set(self.format_currency(usd_rate, "CLP") if usd_rate else "Sin tipo de cambio cargado")
		self.total_expenses_value.set(self.format_currency(total_expenses, "CLP"))
		self.available_balance_value.set(self.format_currency(available_balance, "CLP"))
		self.expense_count_value.set(str(expense_count))
		self.average_expense_value.set(self.format_currency(average_expense, "CLP"))
		self.savings_clp_value.set(self.format_currency(savings_totals["CLP"], "CLP"))
		self.savings_usd_value.set(self.format_currency(savings_totals["USD"], "USD"))
		self.savings_usd_clp_value.set(self.format_currency(savings_totals["USD_TO_CLP"], "CLP"))

		for item in self.expense_tree.get_children():
			self.expense_tree.delete(item)
		for expense in sorted(month_expenses, key=lambda item: (item["date"], item["description"])):
			self.expense_tree.insert(
				"",
				"end",
				iid=expense["id"],
				values=(
					expense["date"],
					expense["description"],
					expense["category"],
					self.format_currency(float(expense["amount"]), "CLP"),
				),
			)

		for item in self.category_tree.get_children():
			self.category_tree.delete(item)
		totals_by_category = self.get_month_totals_by_category()
		for category, total in sorted(totals_by_category.items(), key=lambda item: item[1], reverse=True):
			self.category_tree.insert("", "end", values=(category, self.format_currency(total, "CLP")))

		for item in self.savings_tree.get_children():
			self.savings_tree.delete(item)
		for saving in sorted(self.get_month_savings(), key=lambda item: (item["date"], item["description"])):
			equivalent_clp = float(saving["amount"])
			if saving["currency"] == "USD":
				equivalent_clp = float(saving["amount"]) * usd_rate
			self.savings_tree.insert(
				"",
				"end",
				iid=saving["id"],
				values=(
					saving["date"],
					saving["description"],
					saving.get("goal", "Otro"),
					saving["currency"],
					self.format_currency(float(saving["amount"]), saving["currency"]),
					self.format_currency(equivalent_clp, "CLP"),
				),
			)

		self.draw_category_chart(totals_by_category)
		savings_by_goal = self.get_savings_totals_by_goal()
		self.draw_savings_chart(savings_by_goal)

	def draw_category_chart(self, totals_by_category: dict[str, float]) -> None:
		canvas = self.chart_canvas
		canvas.delete("all")
		canvas.update_idletasks()

		width = max(canvas.winfo_width(), 400)
		height = max(canvas.winfo_height(), 220)
		canvas.create_text(18, 18, anchor="nw", text="📉 Distribución de gastos por categoría en CLP", fill=COLORS["primary"], font=("Segoe UI", 12, "bold"))

		if not totals_by_category:
			canvas.create_text(width / 2, height / 2, text="No hay gastos registrados para este mes.", fill=COLORS["text_secondary"], font=("Segoe UI", 11))
			return

		items = sorted(totals_by_category.items(), key=lambda item: item[1], reverse=True)
		max_total = max(total for _, total in items)
		left = 18
		top = 52
		bar_height = 22
		gap = 12
		max_bar_width = width - 300

		for index, (category, total) in enumerate(items):
			y = top + index * (bar_height + gap)
			color = CHART_COLORS[index % len(CHART_COLORS)]
			bar_width = 0 if max_total == 0 else (total / max_total) * max_bar_width
			emoji = CATEGORY_EMOJIS.get(category, "•")
			canvas.create_text(left, y + bar_height / 2, anchor="w", text=f"{emoji} {category}", fill=COLORS["text_primary"], font=("Segoe UI", 10))
			canvas.create_rectangle(left + 130, y, left + 130 + bar_width, y + bar_height, fill=color, width=0)
			canvas.create_text(
				left + 140 + max_bar_width,
				y + bar_height / 2,
				anchor="e",
				text=self.format_currency(total, "CLP"),
				fill=COLORS["text_primary"],
				font=("Segoe UI", 10, "bold"),
			)

	def draw_savings_chart(self, totals_by_goal: dict[str, float]) -> None:
		canvas = self.savings_chart_canvas
		canvas.delete("all")
		canvas.update_idletasks()

		width = max(canvas.winfo_width(), 400)
		height = max(canvas.winfo_height(), 220)
		canvas.create_text(18, 18, anchor="nw", text="📈 Distribución de ahorros por meta en CLP", fill=COLORS["primary"], font=("Segoe UI", 12, "bold"))

		if not totals_by_goal:
			canvas.create_text(width / 2, height / 2, text="No hay ahorros registrados para este mes.", fill=COLORS["text_secondary"], font=("Segoe UI", 11))
			return

		items = sorted(totals_by_goal.items(), key=lambda item: item[1], reverse=True)
		max_total = max(total for _, total in items)
		left = 18
		top = 52
		bar_height = 22
		gap = 12
		max_bar_width = width - 250

		for index, (goal, total) in enumerate(items):
			y = top + index * (bar_height + gap)
			color = CHART_COLORS[index % len(CHART_COLORS)]
			bar_width = 0 if max_total == 0 else (total / max_total) * max_bar_width
			emoji = GOAL_EMOJIS.get(goal, "•")
			canvas.create_text(left, y + bar_height / 2, anchor="w", text=f"{emoji} {goal}", fill=COLORS["text_primary"], font=("Segoe UI", 10))
			canvas.create_rectangle(left + 120, y, left + 120 + bar_width, y + bar_height, fill=color, width=0)
			canvas.create_text(
				left + 130 + max_bar_width,
				y + bar_height / 2,
				anchor="e",
				text=self.format_currency(total, "CLP"),
				fill=COLORS["text_primary"],
				font=("Segoe UI", 10, "bold"),
			)

	def create_month_report(self, month: str) -> Path:
		month = validate_month(month)
		report_path = _base_dir() / f"reporte_{month}.docx"
		month_expenses = sorted(self.get_month_expenses(month), key=lambda item: (item["date"], item["description"]))
		month_savings = sorted(self.get_month_savings(month), key=lambda item: (item["date"], item["description"]))
		totals_by_category = self.get_month_totals_by_category(month)
		savings_by_goal = self.get_savings_totals_by_goal(month)
		salary = self.get_current_salary(month)
		usd_rate = self.get_current_usd_rate(month)
		total_expenses = sum(float(expense["amount"]) for expense in month_expenses)
		savings_totals = self.get_savings_totals(month)
		balance = salary - total_expenses - savings_totals["TOTAL_CLP_EQUIVALENT"]

		document = Document()
		document.add_heading(f"Reporte mensual de gastos - {month}", level=1)
		document.add_paragraph(f"Fecha de generacion: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

		summary = document.add_table(rows=7, cols=2)
		summary.style = "Light List Accent 1"
		summary.cell(0, 0).text = "Salario en CLP"
		summary.cell(0, 1).text = self.format_currency(salary, "CLP")
		summary.cell(1, 0).text = "Tipo de cambio USD a CLP"
		summary.cell(1, 1).text = self.format_currency(usd_rate, "CLP") if usd_rate else "No registrado"
		summary.cell(2, 0).text = "Total gastado en CLP"
		summary.cell(2, 1).text = self.format_currency(total_expenses, "CLP")
		summary.cell(3, 0).text = "Ahorro en CLP"
		summary.cell(3, 1).text = self.format_currency(savings_totals["CLP"], "CLP")
		summary.cell(4, 0).text = "Ahorro en USD"
		summary.cell(4, 1).text = self.format_currency(savings_totals["USD"], "USD")
		summary.cell(5, 0).text = "Ahorro USD equivalente en CLP"
		summary.cell(5, 1).text = self.format_currency(savings_totals["USD_TO_CLP"], "CLP")
		summary.cell(6, 0).text = "Disponible final en CLP"
		summary.cell(6, 1).text = self.format_currency(balance, "CLP")

		document.add_paragraph("")
		document.add_heading("Gastos del mes", level=2)
		if month_expenses:
			table = document.add_table(rows=1, cols=4)
			table.style = "Table Grid"
			headers = table.rows[0].cells
			headers[0].text = "Fecha"
			headers[1].text = "Descripcion"
			headers[2].text = "Categoria"
			headers[3].text = "Monto CLP"

			for expense in month_expenses:
				row = table.add_row().cells
				row[0].text = expense["date"]
				row[1].text = expense["description"]
				row[2].text = expense["category"]
				row[3].text = self.format_currency(float(expense["amount"]), "CLP")
		else:
			document.add_paragraph("No se registraron gastos para este mes.")

		document.add_paragraph("")
		document.add_heading("Ahorros del mes", level=2)
		if month_savings:
			table = document.add_table(rows=1, cols=6)
			table.style = "Table Grid"
			headers = table.rows[0].cells
			headers[0].text = "Fecha"
			headers[1].text = "Descripcion"
			headers[2].text = "Meta"
			headers[3].text = "Moneda"
			headers[4].text = "Monto"
			headers[5].text = "Equivalente CLP"

			for saving in month_savings:
				equivalent_clp = float(saving["amount"])
				if saving["currency"] == "USD":
					equivalent_clp *= usd_rate
				row = table.add_row().cells
				row[0].text = saving["date"]
				row[1].text = saving["description"]
				row[2].text = saving.get("goal", "Otro")
				row[3].text = saving["currency"]
				row[4].text = self.format_currency(float(saving["amount"]), saving["currency"])
				row[5].text = self.format_currency(equivalent_clp, "CLP")
		else:
			document.add_paragraph("No se registraron ahorros para este mes.")

		document.add_paragraph("")
		document.add_heading("Totales de ahorro por meta en CLP", level=2)
		if savings_by_goal:
			table = document.add_table(rows=1, cols=2)
			table.style = "Table Grid"
			headers = table.rows[0].cells
			headers[0].text = "Meta"
			headers[1].text = "Total equivalente CLP"

			for goal, total in sorted(savings_by_goal.items(), key=lambda item: item[1], reverse=True):
				row = table.add_row().cells
				row[0].text = goal
				row[1].text = self.format_currency(total, "CLP")
		else:
			document.add_paragraph("No hay metas de ahorro con movimientos en este periodo.")

		document.add_paragraph("")
		document.add_heading("Totales de gastos por categoria", level=2)
		if totals_by_category:
			table = document.add_table(rows=1, cols=2)
			table.style = "Table Grid"
			headers = table.rows[0].cells
			headers[0].text = "Categoria"
			headers[1].text = "Total CLP"

			for category, total in sorted(totals_by_category.items(), key=lambda item: item[1], reverse=True):
				row = table.add_row().cells
				row[0].text = category
				row[1].text = self.format_currency(total, "CLP")
		else:
			document.add_paragraph("No hay categorias con gasto en este periodo.")

		document.save(report_path)
		return report_path

	def export_selected_month_report(self) -> None:
		try:
			report_path = self.create_month_report(self.selected_month.get().strip())
		except ValueError:
			messagebox.showerror("Mes invalido", "No se pudo generar el Word porque el mes no es valido.")
			return

		self.set_status(f"Reporte exportado: {report_path.name}")
		messagebox.showinfo("Reporte creado", f"Se genero el archivo {report_path.name} en la carpeta del proyecto.")

	def handle_system_month_change(self, show_message: bool) -> None:
		current_system_month = date.today().strftime("%Y-%m")
		metadata = self.data.setdefault("metadata", {})
		previous_system_month = metadata.get("last_seen_system_month", current_system_month)

		if previous_system_month != current_system_month:
			report_path = None
			if month_has_activity(self.data, previous_system_month):
				report_path = self.create_month_report(previous_system_month)

			metadata["last_seen_system_month"] = current_system_month
			self.save_data()
			self.selected_month.set(current_system_month)
			self.reset_expense_form()
			self.reset_saving_form()

			if report_path is not None:
				message = f"Cambio de mes detectado. Se genero {report_path.name} con el movimiento de {previous_system_month}."
			else:
				message = f"Cambio de mes detectado. No habia movimientos para exportar en {previous_system_month}."

			self.set_status(message)
			if show_message:
				messagebox.showinfo("Nuevo mes detectado", message)
		else:
			metadata["last_seen_system_month"] = current_system_month
			self.save_data()

	def schedule_system_month_check(self) -> None:
		self.root.after(SYSTEM_MONTH_CHECK_MS, self.check_system_month_rollover)

	def check_system_month_rollover(self) -> None:
		self.handle_system_month_change(show_message=True)
		self.refresh_view()
		self.schedule_system_month_check()


def main() -> None:
	root = tk.Tk()
	app = ExpenseApp(root)
	root.mainloop()


if __name__ == "__main__":
	main()